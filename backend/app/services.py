import io
import json
import re
import unicodedata
import zipfile
from urllib.parse import urlsplit

import httpx
from fastapi import HTTPException
from sqlalchemy import select

from .database import Document, Search


def words(text: str) -> set[str]:
    plain = unicodedata.normalize("NFKD", text.casefold())
    plain = "".join(c for c in plain if not unicodedata.combining(c))
    return {x for x in re.findall(r"\w+", plain) if len(x) > 2}


def search_documents(db, query: str):
    terms = words(query)
    results = []
    for doc in db.scalars(select(Document)):
        # Fragmentos locais; nenhum documento é enviado ao provedor nesta busca.
        for offset in range(0, len(doc.content), 700):
            chunk = doc.content[offset:offset + 1000]
            score = len(terms & words(chunk))
            if score:
                results.append({"document_id": doc.id, "name": doc.name, "excerpt": chunk,
                                "offset": offset, "score": score})
    return sorted(results, key=lambda r: r["score"], reverse=True)[:6]


def extract_document(name: str, raw: bytes) -> str:
    extension = name.rsplit(".", 1)[-1].lower()
    try:
        if extension == "txt":
            text = raw.decode("utf-8-sig")
        elif extension == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 200:
                raise ValueError("PDF criptografado ou acima de 200 páginas.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == "docx":
            from docx import Document as WordDocument
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30_000_000:
                    raise ValueError("Documento expandido muito grande.")
            doc = WordDocument(io.BytesIO(raw))
            text = "\n".join([p.text for p in doc.paragraphs] +
                             [" ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        else:
            raise HTTPException(415, "Use PDF, TXT UTF-8 ou DOCX.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Não foi possível extrair texto. Use arquivo válido, sem senha, até 200 páginas.")
    if not text.strip():
        raise HTTPException(422, "Documento sem texto extraível. PDFs digitalizados precisam de OCR externo.")
    if len(text) > 500_000:
        raise HTTPException(413, "O texto excede o limite de 500 mil caracteres.")
    return text


async def web_search(settings, db, query: str):
    if not settings.search_key:
        raise HTTPException(503, "Busca não configurada. Adicione JARVIS_SEARCH_KEY no backend (Brave Search).")
    try:
        async with httpx.AsyncClient(timeout=20, follow_redirects=False) as client:
            response = await client.get("https://api.search.brave.com/res/v1/web/search",
                params={"q": query, "count": 5, "search_lang": "pt-br", "country": "BR"},
                headers={"X-Subscription-Token": settings.search_key, "Accept": "application/json"})
            response.raise_for_status()
            payload = response.json()
        results = []
        for r in payload.get("web", {}).get("results", [])[:5]:
            url = r.get("url", "")
            parsed = urlsplit(url)
            if parsed.scheme not in ("http", "https") or not parsed.hostname or parsed.username:
                continue
            results.append({"title": r.get("title", "Sem título"), "url": url,
                            "description": r.get("description", ""), "age": r.get("age")})
        db.add(Search(query=query, result=json.dumps(results, ensure_ascii=False)))
        db.commit()
        return {"query": query, "provider": "Brave Search", "results": results}
    except (httpx.HTTPError, ValueError, TypeError):
        raise HTTPException(502, "O provedor de busca falhou. Verifique acesso, chave e cota no backend.")


async def send_telegram(settings, recipient: str, text: str) -> str:
    if not settings.telegram_token:
        raise HTTPException(503, "Telegram não configurado no backend.")
    try:
        async with httpx.AsyncClient(timeout=20, follow_redirects=False) as client:
            response = await client.post(
                f"https://api.telegram.org/bot{settings.telegram_token}/sendMessage",
                json={"chat_id": recipient, "text": text})
            payload = response.json()
            if response.status_code != 200 or not payload.get("ok"):
                raise HTTPException(502, "Telegram recusou o envio. Verifique chat_id, bot e permissões.")
            return str(payload["result"]["message_id"])
    except (httpx.HTTPError, ValueError, KeyError):
        # Uma falha de transporte pode ocorrer depois de o provedor receber a mensagem.
        raise HTTPException(502, "Entrega incerta: consulte o Telegram antes de criar outro rascunho.")
